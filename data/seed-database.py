"""种子词典生成 + 500行病历灌入 tcm_ehr.records
五实体术语库：疾病(disease)/证候(pattern)/症状(symptom)/中药(herb)/方剂(formula)
"""
import openpyxl, pymysql, re, json, uuid, random, collections

XLSX = r'D:\ZISHIKU\AI\tcm-ehr-governance\docs\电子病历精简脱敏数据_500行.xlsx'
DICT_DIR = r'D:\ZISHIKU\AI\tcm-ehr-governance\data\dictionaries'

wb = openpyxl.load_workbook(XLSX)
ws = wb.active
headers = [str(c.value).strip() if c.value else '' for c in ws[1]]
I = {h: i for i, h in enumerate(headers)}

rows = []
for r in ws.iter_rows(min_row=2, values_only=True):
    if not r[I['登记号']]:
        continue
    rows.append(r)
print(f'Excel有效行数: {len(rows)}')

# ============ 1. 提取五类实体 ============

# 疾病：中医诊断列 "胃痛-肝胃不和证" → 提取疾病名 "胃痛"
disease_counter = collections.Counter()
for r in rows:
    v = str(r[I['中医诊断']] or '')
    # 格式：疾病名-证候名，取"-"前的疾病名
    if v:
        disease = v.split('-')[0].strip() if '-' in v else v.strip()
        if disease and len(disease) >= 2:
            disease_counter[disease] += 1

# 证候：辨证结论列
pattern_set = []
seen_pattern = set()
for r in rows:
    v = r[I['辨证结论']]
    if v:
        for part in re.split(r'[、，,；;]', str(v)):
            p = part.strip()
            if p and p not in seen_pattern:
                seen_pattern.add(p)
                pattern_set.append(p)

# 症状：主诉列（口语化→标准术语）
def clean_symptom(s):
    s = re.sub(r'\d+(\.\d+)?\s*(月|天|周|年|小时|日)$', '', s.strip())
    return s.strip()

symptom_counter = collections.Counter()
for r in rows:
    v = r[I['主诉']]
    if v:
        for part in re.split(r'[、，,；;]', str(v)):
            s = clean_symptom(part)
            if s and 1 < len(s) <= 8:
                symptom_counter[s] += 1

# 中药：草药列
herb_counter = collections.Counter()
for r in rows:
    v = str(r[I['草药']] or '')
    v = re.sub(r'^中草药处方\(\d+付\)[，,]?', '', v)
    for part in re.split(r'[，,]', v):
        m = re.match(r'^([\u4e00-\u9fa5]{1,6})\s*\d+(\.\d+)?\s*(g|片|枚|条|张|付|ml|mg)?$', part.strip())
        if m:
            herb_counter[m.group(1)] += 1

# 方剂：按核心药组合推断
FORMULA_RULES = [
    ('柴胡疏肝散', ['柴胡', '白芍', '枳壳', '香附']),
    ('独活寄生汤', ['独活', '桑寄生', '杜仲']),
    ('知柏地黄丸', ['熟地黄', '山茱萸', '知母', '黄柏']),
    ('济生肾气丸', ['附子', '肉桂', '牛膝', '车前子']),
    ('白虎加人参汤', ['石膏', '知母', '人参', '粳米']),
    ('真武汤', ['附子', '白术', '茯苓', '白芍']),
    ('七味白术散', ['人参', '白术', '茯苓', '葛根']),
]
formula_alias = {
    '柴胡疏肝散': ['疏肝散', '柴胡疏肝汤'],
    '独活寄生汤': ['寄生汤'],
    '知柏地黄丸': ['知柏地黄汤'],
    '济生肾气丸': ['肾气丸'],
    '白虎加人参汤': ['白虎汤加人参'],
    '真武汤': [],
    '七味白术散': ['七味白术汤'],
}
formula_counter = collections.Counter()
for r in rows:
    v = str(r[I['草药']] or '')
    herbs_in_rx = set(re.findall(r'([\u4e00-\u9fa5]{1,6})\d+(?:\.\d+)?g', v))
    for name, core in FORMULA_RULES:
        if all(c in herbs_in_rx for c in core):
            formula_counter[name] += 1
            break

print(f'疾病: {len(disease_counter)} | 证候: {len(pattern_set)} | 症状: {len(symptom_counter)} | 中药: {len(herb_counter)} | 方剂: {len(formula_counter)}')

# ============ 2. 生成词典JSON ============
def write_json(fname, data):
    path = f'{DICT_DIR}\\{fname}'
    with open(path, 'w', encoding='utf-8') as f:
        json.dump(data, f, ensure_ascii=False, indent=1)
    print(f'  写入 {fname}: {len(data)}条')

diseases = [{'standardTerm': d, 'aliases': [], 'source': '中医临床诊疗术语 疾病'}
            for d, c in disease_counter.most_common()]

patterns = [{'standardTerm': s, 'aliases': [], 'source': '中医病证分类与代码 GB/T 15657-2021'}
            for s in pattern_set]

symptoms = []
for s, c in symptom_counter.most_common():
    if c < 3:
        break
    aliases = []
    if '痛' in s and len(s) >= 3:
        aliases.append(s.replace('痛', '疼痛'))
    if '胀' in s:
        aliases.append(s.replace('胀', '胀满'))
    symptoms.append({'standardTerm': s, 'aliases': aliases, 'source': '中医临床诊疗术语 症状'})

herbs = [{'standardTerm': h, 'aliases': [], 'source': '中国药典2025年版'}
         for h, c in herb_counter.most_common()]

formulas = [{'standardTerm': n, 'aliases': formula_alias.get(n, []), 'source': '中医方剂大辞典'}
            for n, c in formula_counter.most_common()]

write_json('diseases.json', diseases)
write_json('patterns.json', patterns)
write_json('symptoms.json', symptoms)
write_json('herbs.json', herbs)
write_json('formulas.json', formulas)

# ============ 3. 灌入records表 ============
conn = pymysql.connect(host='localhost', port=3306, user='root', password='123456',
                       database='tcm_ehr', charset='utf8mb4')
cur = conn.cursor()
cur.execute('SET FOREIGN_KEY_CHECKS=0')
cur.execute('TRUNCATE TABLE records')
cur.execute('SET FOREIGN_KEY_CHECKS=1')
print('records表已清空，重新灌入')

herb_names = set(herb_counter.keys())
symptom_names = set(symptom_counter.keys())
inserted = 0

for r in rows:
    rid = str(uuid.uuid4())
    chief = str(r[I['主诉']] or '')
    self_report = str(r[I['自诉']] or '')
    tongue = str(r[I['舌诊']] or '').strip()
    pulse = str(r[I['脉诊']] or '').strip()

    # diseases: 从中医诊断提取疾病名
    tcm_diag = str(r[I['中医诊断']] or '')
    disease_name = tcm_diag.split('-')[0].strip() if '-' in tcm_diag else tcm_diag.strip()
    disease_entities = [{'content': disease_name, 'sourceText': tcm_diag}] if disease_name else []

    # symptoms: 主诉/自诉提取症状词
    sym_entities = []
    seen_sym = set()
    for frag in re.split(r'[、，,；;]', chief):
        frag = frag.strip()
        for s in symptom_names:
            if s in frag and s not in seen_sym:
                seen_sym.add(s)
                sym_entities.append({'content': s, 'sourceText': frag})
    for s in symptom_names:
        if s in self_report and s not in seen_sym:
            seen_sym.add(s)
            sym_entities.append({'content': s, 'sourceText': self_report[:30]})

    # patternList: 证候实体
    pattern_str = str(r[I['辨证结论']] or '').strip()
    pattern_entities = [{'content': p.strip(), 'sourceText': p.strip()}
                       for p in re.split(r'[、，,；;]', pattern_str) if p.strip()]

    # tongueList / pulseList
    tongue_entities = [{'content': tongue, 'sourceText': tongue}] if tongue else []
    pulse_entities = [{'content': pulse, 'sourceText': pulse}] if pulse else []

    # herbs: 药名+剂量
    v_herb = str(r[I['草药']] or '')
    herb_entities = []
    for m in re.finditer(r'([\u4e00-\u9fa5]{1,6})(\d+(?:\.\d+)?(?:g|片|枚|条|张|付|ml|mg)?)', v_herb):
        name = m.group(1)
        if name in herb_names:
            herb_entities.append({'name': name, 'dosage': m.group(2), 'sourceText': m.group(0)})

    # formulaList: 核心药组合推断
    formula = ''
    herbs_in_rx = set(h['name'] for h in herb_entities)
    for name, core in FORMULA_RULES:
        if all(c in herbs_in_rx for c in core):
            formula = name
            break
    formula_entities = [{'content': formula, 'sourceText': v_herb[:30]}] if formula else []

    structured = json.dumps({
        'diseases': disease_entities,
        'symptoms': sym_entities,
        'tongueList': tongue_entities,
        'pulseList': pulse_entities,
        'patternList': pattern_entities,
        'causeList': [],
        'treatmentList': [],
        'formulaList': formula_entities,
        'herbs': herb_entities
    }, ensure_ascii=False)

    roll = random.random()
    if roll < 0.85:
        score, grade, status = random.randint(80, 96), '合格', 'completed'
    elif roll < 0.97:
        score, grade, status = random.randint(70, 79), '待复核', 'pending'
    else:
        score, grade, status = random.randint(40, 60), '无效', 'invalid'

    visit_time = str(r[I['接诊时间']] or '')[:19]
    sql = '''INSERT INTO records (id, registration_no, outpatient_no, gender, age, visit_count,
        western_diagnosis, tcm_diagnosis, present_illness, chief_complaint, self_report,
        inspection, pulse, tongue, physical_exam, pattern, prescription, follow_up,
        treatment_effect, department, doctor_id, visit_time, structured_data, score, grade, status)
        VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)'''
    vals = (rid, str(r[I['登记号']]), str(r[I['门诊号']] or ''), str(r[I['性别']] or ''),
            str(r[I['年龄']] or ''), int(r[I['就诊次数']] or 1),
            str(r[I['西医诊断']] or ''), str(r[I['中医诊断']] or ''), str(r[I['现病史']] or ''),
            str(r[I['主诉']] or ''), str(r[I['自诉']] or ''), str(r[I['望诊']] or ''),
            str(r[I['脉诊']] or ''), str(r[I['舌诊']] or ''), str(r[I['查体']] or ''),
            str(r[I['辨证结论']] or ''), str(r[I['草药']] or ''), str(r[I['随访']] or ''),
            str(r[I['治疗效果']] or ''), str(r[I['开单科室']] or '中医内科'),
            str(r[I['医生工号']] or ''), visit_time if visit_time else None,
            structured, score, grade, status)
    cur.execute(sql, vals)
    inserted += 1

conn.commit()
print(f'灌入病历: {inserted}条')

cur.execute("SELECT grade, COUNT(*) FROM records GROUP BY grade")
for g, c in cur.fetchall():
    print(f'  {g}: {c}')
conn.close()
print('DONE')
